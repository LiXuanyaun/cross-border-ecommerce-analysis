"""Export one analysis bundle to Excel, Markdown, DOCX, and a manifest."""
from datetime import datetime
from math import ceil
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any, Dict, Iterable, List
import json
import textwrap

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from autoclean.analytics import AnalysisStatus, DocxReportBuilder, ExcelBundleWriter
from crossborder_analytics.localization import (
    MODULE_LABELS,
    STATUS_LABELS,
    currency_label,
    field_guide_frame,
    localize_frame,
    value_for,
)


COLORS = {
    "ink": "#17212B",
    "paper": "#F7F8F6",
    "grid": "#DDE3E0",
    "teal": "#0D7C66",
    "coral": "#C84B31",
    "amber": "#B98213",
}

REPORT_FONT_CANDIDATES = (
    "Microsoft YaHei", "PingFang SC", "Noto Sans CJK SC", "Noto Sans SC",
    "Source Han Sans SC", "SimHei", "Arial Unicode MS", "DejaVu Sans",
)


def _available_report_font() -> str:
    for family in REPORT_FONT_CANDIDATES:
        try:
            font_manager.findfont(family, fallback_to_default=False)
            return family
        except ValueError:
            continue
    return "DejaVu Sans"


REPORT_FONT_FAMILY = _available_report_font()
matplotlib.rcParams["font.sans-serif"] = [REPORT_FONT_FAMILY, *REPORT_FONT_CANDIDATES]
matplotlib.rcParams["axes.unicode_minus"] = False


def _wrap_report_label(value: Any, width: int = 10) -> str:
    return textwrap.fill(str(value), width=width, break_long_words=True, break_on_hyphens=False)


def _report_figure(figsize):
    return plt.subplots(figsize=figsize, constrained_layout=True)


def _style_report_axis(ax, title: str, *, xlabel: str = "", ylabel: str = "", grid_axis: str = "") -> None:
    ax.set_title(title, loc="left", fontweight="bold", fontsize=12, pad=14, color=COLORS["ink"])
    if xlabel:
        ax.set_xlabel(xlabel, labelpad=10, color=COLORS["ink"])
    if ylabel:
        ax.set_ylabel(ylabel, labelpad=12, color=COLORS["ink"])
    ax.tick_params(axis="both", colors=COLORS["ink"], labelsize=9, pad=5)
    if grid_axis:
        ax.grid(axis=grid_axis, color=COLORS["grid"], linewidth=.8)
    ax.set_axisbelow(True)
    ax.spines[["top", "right"]].set_visible(False)
    ax.margins(x=.035, y=.08)


def _save_report_chart(fig, path: Path) -> None:
    fig.savefig(
        path,
        dpi=180,
        bbox_inches="tight",
        pad_inches=.24,
        facecolor="white",
    )
    plt.close(fig)


def _result(bundle, name):
    result = bundle.results.get(name)
    return result if result and result.status == AnalysisStatus.SUCCESS else None


def _frame_or_empty(value, columns=None):
    if isinstance(value, pd.DataFrame):
        return value.copy()
    return pd.DataFrame(value or [], columns=columns)


def _flat_frame(frame: pd.DataFrame) -> pd.DataFrame:
    """Convert nested artifact fields to stable text for tabular exporters."""
    if not isinstance(frame, pd.DataFrame):
        return pd.DataFrame(frame)
    output = frame.copy()
    for column in output.columns:
        if output[column].map(lambda value: isinstance(value, (dict, list, tuple, set))).any():
            def serialize(value):
                if not isinstance(value, (dict, list, tuple, set)):
                    return value
                text = json.dumps(value, ensure_ascii=False, sort_keys=True)
                return text if len(text) <= 32000 else text[:31960] + "... [预览已截断]"
            output[column] = output[column].map(serialize)
    return output


def _clip_report_text(value: Any, limit: int = 180) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    if isinstance(value, (dict, list, tuple, set)):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True)
    text = str(value)
    return text if len(text) <= limit else text[:limit - 3] + "..."


def _docx_evidence_bullets(frame: pd.DataFrame) -> list[str]:
    """Keep the DOCX evidence appendix readable; full objects stay in Excel/manifest."""
    if frame.empty:
        return []
    items = []
    for row in frame.head(40).to_dict("records"):
        items.append(
            "证据 {}；查询 {}；期间 {} 至 {}；样本 {:,} 行；公式：{}；结果摘要：{}；限制：{}。".format(
                _clip_report_text(row.get("evidence_id"), 36),
                _clip_report_text(row.get("query_name"), 48),
                _clip_report_text(row.get("period_start"), 20),
                _clip_report_text(row.get("period_end"), 20),
                int(row.get("row_count") or 0),
                _clip_report_text(row.get("formula"), 160),
                _clip_report_text(row.get("result_digest"), 24),
                _clip_report_text(row.get("limitations"), 140),
            )
        )
    return items


def _scope_frames(artifacts, report_scope="overall", scope_value=None):
    markets = pd.DataFrame(artifacts.records("market_opportunities")) if artifacts else pd.DataFrame()
    products = pd.DataFrame(artifacts.records("product_opportunities")) if artifacts else pd.DataFrame()
    actions = pd.DataFrame(artifacts.records("action_items")) if artifacts else pd.DataFrame()
    if report_scope == "market" and scope_value:
        markets = markets.loc[markets.market.astype(str).eq(str(scope_value))] if not markets.empty else markets
        if not products.empty:
            products = products.loc[
                products.primary_market.astype(str).eq(str(scope_value))
                | products.target_market.astype(str).eq(str(scope_value))
            ]
    elif report_scope == "category" and scope_value:
        products = products.loc[products.category.astype(str).eq(str(scope_value))] if not products.empty else products
        markets = markets.iloc[0:0]
    if not actions.empty and report_scope != "overall":
        source_ids = set(markets.get("opportunity_id", pd.Series(dtype=str)).astype(str)) | set(products.get("opportunity_id", pd.Series(dtype=str)).astype(str))
        actions = actions.loc[actions.source_id.astype(str).isin(source_ids)]
    return markets, products, actions


def _business_report_frame(payload: dict[str, Any]) -> pd.DataFrame:
    rows = []
    for metric in payload.get("metrics", []):
        rows.append({"record_type": "metric", **metric})
    for anomaly in payload.get("anomalies", []):
        rows.append({"record_type": "anomaly", **anomaly})
    for action in payload.get("actions", []):
        rows.append({"record_type": "action", **action})
    return pd.DataFrame(rows)


def _style_excel_report(path: Path) -> None:
    """Make report tables legible regardless of the spreadsheet viewer theme."""
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="17212B")
    body_fill = PatternFill("solid", fgColor="F7F8F6")
    header_font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
    body_font = Font(name="Microsoft YaHei", color="17212B")
    header_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    body_alignment = Alignment(vertical="top", wrap_text=True)
    border = Border(bottom=Side(style="hair", color="DDE3E0"))
    for worksheet in workbook.worksheets:
        worksheet.sheet_view.showGridLines = False
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.calculate_dimension()
        for row_index, row in enumerate(worksheet.iter_rows(), start=1):
            for cell in row:
                cell.fill = header_fill if row_index == 1 else body_fill
                cell.font = header_font if row_index == 1 else body_font
                cell.alignment = header_alignment if row_index == 1 else body_alignment
                cell.border = border
        worksheet.row_dimensions[1].height = 28
    workbook.save(path)


def output_frames(bundle, report_scope="overall", scope_value=None, business_analysis=None) -> Dict[str, pd.DataFrame]:
    overview = _result(bundle, "overview")
    summary = pd.DataFrame([
        {"metric": key, "value": value, "currency": bundle.context.metadata.get("target_currency", "")}
        for key, value in (overview.data.items() if overview else [])
    ])
    product = _result(bundle, "product")
    customer = _result(bundle, "customer")
    region = _result(bundle, "region")
    market_category = _result(bundle, "market_category")
    returns = _result(bundle, "returns")

    return_parts = []
    if returns:
        for dimension in ("category", "region", "products"):
            frame = returns.data.get(dimension)
            if isinstance(frame, pd.DataFrame) and not frame.empty:
                part = frame.copy()
                part.insert(0, "analysis_type", dimension)
                return_parts.append(part)
        return_parts.append(pd.DataFrame([{"analysis_type": "summary", **returns.data["summary"]}]))

    fx_rates = bundle.metadata.get("fx_rates")
    if not isinstance(fx_rates, pd.DataFrame):
        fx_rates = pd.DataFrame()
    artifacts = getattr(bundle, "artifacts", None)
    phase2_evidence = pd.DataFrame(artifacts.records("evidence")) if artifacts else pd.DataFrame()
    raw_frames = {
        "summary": summary,
        "product_analysis": product.data["products"] if product else pd.DataFrame(),
        "customer_analysis": customer.data["customers"] if customer else pd.DataFrame(),
        "customer_composition": customer.data.get("composition", pd.DataFrame()) if customer else pd.DataFrame(),
        "region_analysis": region.data if region else pd.DataFrame(),
        "market_category": market_category.data if market_category else pd.DataFrame(),
        "product_detail_sources": pd.DataFrame([
            {"name": query, "status": "SUCCESS", "message": "按dataset_id、分析周期和product_id参数化查询"}
            for query in bundle.metadata.get("product_detail_queries", [])
        ]),
        "return_analysis": pd.concat(return_parts, ignore_index=True, sort=False) if return_parts else pd.DataFrame(),
        "module_status": bundle.status_frame(),
        "evidence": phase2_evidence if not phase2_evidence.empty else bundle.evidence_frame(),
        "legacy_evidence": bundle.evidence_frame(),
        "data_quality": pd.DataFrame([issue.to_dict() for issue in bundle.context.issues]),
        "fx_rates": fx_rates,
    }
    if artifacts:
        market_opportunities, product_opportunities, action_items = _scope_frames(
            artifacts, report_scope, scope_value,
        )
        raw_frames.update({
            "metric_definitions": pd.DataFrame(artifacts.records("metric_definitions")),
            "metric_snapshots": pd.DataFrame(artifacts.records("metric_snapshots")),
            "anomalies": pd.DataFrame(artifacts.records("anomalies")),
            "diagnoses": pd.DataFrame(artifacts.records("diagnoses")),
            "recommendations": pd.DataFrame(artifacts.records("recommendations")),
            "data_quality_summary": pd.DataFrame(artifacts.records("data_quality")),
            "data_quality_dimensions": pd.DataFrame(artifacts.records("data_quality_dimensions")),
            "field_quality": pd.DataFrame(artifacts.records("field_quality")),
            "analysis_capability": pd.DataFrame(artifacts.records("analysis_capability")),
            "data_improvement_plan": pd.DataFrame(artifacts.records("data_improvement_plan")),
            "data_quality_issues": pd.DataFrame([issue.to_dict() for issue in bundle.context.issues]),
            "insights": pd.DataFrame(artifacts.records("insights")),
            "市场机会": market_opportunities,
            "产品机会": product_opportunities,
            "行动清单": action_items,
            "机会汇总": pd.DataFrame(artifacts.records("opportunity_summary")),
        })
    if business_analysis:
        labels = {"advertising": "广告分析", "returns": "退款分析", "logistics": "物流分析"}
        raw_frames.update({
            labels[topic]: _business_report_frame(payload)
            for topic, payload in business_analysis.items()
            if topic in labels
        })
        raw_frames["多业务证据"] = pd.DataFrame([
            {"topic": topic, **item}
            for topic, payload in business_analysis.items()
            for item in payload.get("evidence", [])
        ])
    localized = {name: localize_frame(_flat_frame(frame)) for name, frame in raw_frames.items()}
    localized["字段说明"] = field_guide_frame()
    return localized


def _money(value, currency):
    if value is None or pd.isna(value):
        return "无法计算"
    return "{:,.2f} {}".format(float(value), currency_label(currency)).strip()


def _pct(value):
    if value is None or pd.isna(value):
        return "无法计算"
    return "{:.2%}".format(float(value))


def build_markdown(bundle, business_analysis=None) -> str:
    currency = bundle.context.metadata.get("target_currency", "")
    overview = _result(bundle, "overview")
    sales = _result(bundle, "sales")
    product = _result(bundle, "product")
    customer = _result(bundle, "customer")
    region = _result(bundle, "region")
    returns = _result(bundle, "returns")
    lines = [
        "# CrossBorder AI Analytics 经营分析报告",
        "",
        "## 执行摘要",
    ]
    if overview:
        data = overview.data
        lines.extend([
            "- GMV（成交总额）：{} `[overview.gmv]`".format(_money(data["gmv"], currency)),
            "- 订单：{:,}，客户：{:,}，AOV（平均客单价）：{} `[overview.orders]` `[overview.aov]`".format(
                data["orders"], data.get("customers") or 0, _money(data["aov"], currency)
            ),
            "- 利润：{}，利润率：{} `[overview.profit]` `[overview.profit_rate]`".format(
                _money(data.get("profit_amount"), currency), _pct(data.get("profit_rate"))
            ),
            "- 退货率：{} `[overview.return_rate]`".format(_pct(data.get("return_rate"))),
        ])
    lines.extend(["", "## 数据概览"])
    lines.append("- 来源文件：`{}`".format(bundle.context.metadata.get("filename", "")))
    lines.append("- 原始数据哈希：`{}`".format(bundle.context.metadata.get("sha256", "")))
    lines.append("- 分析行数：{:,}".format(len(bundle.context.analysis_data)))
    lines.append("- 币种口径：{} -> {}；汇率来源：{}".format(
        currency_label(bundle.context.metadata.get("source_currency", "")), currency_label(currency), bundle.context.metadata.get("fx_provider", "")
    ))

    lines.extend(["", "## 销售分析"])
    if sales:
        lines.append("最近完整月为 {}，GMV（成交总额）环比 {}。`[sales.latest_complete_gmv]`".format(
            sales.data["latest_complete_month"], _pct(sales.data["mom"])
        ))
        if sales.data["incomplete_months"]:
            lines.append("不完整月份 {} 仅展示，不参与环比结论。".format("、".join(sales.data["incomplete_months"])))
        if not sales.data["seasonality_available"]:
            lines.append("正式季节性分析已跳过：完整月份少于24个。")
    else:
        lines.append(_module_omission(bundle, "sales"))

    lines.extend(["", "## 商品分析"])
    if product:
        counts = product.data["products"].classification.value_counts()
        lines.append("共有 {:,} 个商品达到至少3笔订单的四象限门槛。`[product.matrix]`".format(product.data["eligible_products"]))
        lines.extend("- {}：{:,}".format(label, int(count)) for label, count in counts.items())
    else:
        lines.append(_module_omission(bundle, "product"))

    lines.extend(["", "## 客户分析"])
    if customer:
        lines.append("RFM（客户价值模型）锚点日期：{}。`[customer.rfm]`".format(customer.data["anchor_date"]))
        for row in customer.data["segments"].itertuples():
            lines.append("- {}：{:,}人，GMV（成交总额）{}".format(row.segment, int(row.customers), _money(row.gmv, currency)))
    else:
        lines.append(_module_omission(bundle, "customer"))

    lines.extend(["", "## 市场分析"])
    if region:
        for row in region.data.head(5).itertuples():
            market = value_for(row.market, "region") if row.market_source == "region" else row.market
            lines.append("- {}：GMV（成交总额）{}，订单 {:,}，利润率 {}".format(
                market, _money(row.gmv, currency), int(row.orders), _pct(row.profit_rate)
            ))
    else:
        lines.append(_module_omission(bundle, "region"))

    lines.extend(["", "## 退货与运营风险"])
    if returns:
        lines.append("退货率 {}，退货关联GMV（成交总额）{}。后者不是实际退款损失。`[returns.rate]` `[returns.gmv_exposure]`".format(
            _pct(returns.data["summary"]["return_rate"]),
            _money(returns.data["summary"]["returned_gmv_exposure"], currency),
        ))
        if returns.message:
            lines.append(returns.message + "。")
    else:
        lines.append(_module_omission(bundle, "returns"))

    lines.extend(["", "## 商业建议"])
    if bundle.recommendations:
        for index, item in enumerate(bundle.recommendations, 1):
            lines.append("{}. **{}** {} 依据：{}".format(
                index, item["title"], item["action"], ", ".join("`[{}]`".format(value) for value in item["evidence_ids"])
            ))
    else:
        lines.append("当前证据未达到自动建议阈值，保持监测。")
    artifacts = getattr(bundle, "artifacts", None)
    lines.extend(["", "## 经营异常与行动建议"])
    if artifacts and artifacts.insights:
        for item in artifacts.insights[:10]:
            lines.append("- [{}] {}：{}；建议：{}。`[{}]`".format(
                item.priority, item.entity_name, item.finding, item.recommendation_summary,
                item.insight_id,
            ))
    else:
        lines.append("当前未发现达到规则阈值的异常。")
    lines.extend(["", "## 数据质量与经营分析边界"])
    if artifacts and artifacts.data_quality:
        quality = artifacts.data_quality
        lines.append("- 数据质量评分：{:.2f}/100；可信度等级：{}。{}。".format(
            quality.quality_score, quality.quality_rating, quality.quality_explanation
        ))
        for item in artifacts.unsupported_conclusions[:10]:
            lines.append("- 当前不能判断：{}；原因：{}。".format(
                item["conclusion"], "、".join(item.get("reasons", [])) or "关键证据不足"
            ))
    if business_analysis:
        labels = {"advertising": "广告分析", "returns": "退款分析", "logistics": "物流分析"}
        for topic in ("advertising", "returns", "logistics"):
            payload = business_analysis.get(topic)
            if not payload:
                continue
            lines.extend(["", "## {}（模拟数据）".format(labels[topic])])
            for metric in payload.get("metrics", []):
                lines.append("- {}：{}；公式：`{}`；证据：`[{}]`。".format(
                    metric["label"], metric.get("value"), metric["formula"], metric["evidence_id"],
                ))
            for anomaly in payload.get("anomalies", [])[:5]:
                lines.append("- 异常：{}（{}）；阈值：{}；建议：{}。".format(
                    anomaly["title"], anomaly["entity"], anomaly["threshold"], anomaly["recommendation"],
                ))
        lines.extend(["", "> 广告、退款和物流来自 synthetic_extension 模拟数据，不代表真实经营表现。"])
    lines.extend([
        "", "## 结论", "所有结论均来自本次分析证据对象；缺少字段、样本不足和不完整周期均未被强行推断。",
        "", "## 模块状态", "| 模块 | 状态 | 说明 |", "|---|---|---|",
    ])
    for result in bundle.results.values():
        lines.append("| {} | {} | {} |".format(
            MODULE_LABELS.get(result.name, result.name),
            STATUS_LABELS.get(result.status.value, result.status.value),
            result.message or "-",
        ))
    return "\n".join(lines) + "\n"


def _module_omission(bundle, name):
    result = bundle.results.get(name)
    return "无法分析：{}".format(result.message if result else "模块不存在")


def _save_charts(bundle, directory: Path) -> Dict[str, Path]:
    directory.mkdir(parents=True, exist_ok=True)
    charts = {}
    sales = _result(bundle, "sales")
    if sales and not sales.data["monthly"].empty:
        data = sales.data["monthly"]
        fig, ax = _report_figure((10, 4.8))
        positions = np.arange(len(data))
        labels = data.month.astype(str).tolist()
        ax.plot(positions, data.gmv, color=COLORS["teal"], linewidth=2.2, marker="o", markersize=3)
        stride = max(1, ceil(len(labels) / 12)) if len(labels) > 24 else 1
        tick_indexes = list(range(0, len(labels), stride))
        if tick_indexes and tick_indexes[-1] != len(labels) - 1:
            tick_indexes.append(len(labels) - 1)
        ax.set_xticks(tick_indexes, [labels[index] for index in tick_indexes])
        ax.tick_params(axis="x", rotation=45)
        for label in ax.get_xticklabels():
            label.set_horizontalalignment("right")
        _style_report_axis(ax, "月度GMV（成交总额）趋势", xlabel="月份", ylabel="GMV（成交总额）", grid_axis="y")
        path = directory / "monthly_gmv.png"
        _save_report_chart(fig, path)
        charts["sales"] = path

    product = _result(bundle, "product")
    if product:
        data = product.data["products"]
        eligible = data[data.classification.ne("样本不足")].copy()
        if not eligible.empty:
            fig, ax = _report_figure((9, 5.2))
            palette = {"核心商品": COLORS["teal"], "引流商品": COLORS["amber"], "潜力商品": "#3D6FA3", "淘汰观察": COLORS["coral"]}
            for label, group in eligible.groupby("classification"):
                ax.scatter(group.units, group.profit_rate, s=np.clip(group.gmv / max(group.gmv.max(), 1) * 100, 12, 100), alpha=.58, label=label, color=palette.get(label))
            ax.axvline(product.data["volume_threshold"], color=COLORS["grid"], linewidth=1)
            ax.axhline(product.data["portfolio_margin"], color=COLORS["grid"], linewidth=1)
            _style_report_axis(ax, "商品经营矩阵", xlabel="销量", ylabel="利润率", grid_axis="both")
            ax.legend(
                frameon=False, ncol=2, loc="upper right",
                prop={"family": REPORT_FONT_FAMILY, "size": 9},
            )
            path = directory / "product_matrix.png"
            _save_report_chart(fig, path)
            charts["product"] = path

    region = _result(bundle, "region")
    if region and not region.data.empty:
        data = region.data.sort_values("gmv").tail(10).copy()
        data["market"] = data.apply(
            lambda row: value_for(row["market"], "region") if row["market_source"] == "region" else row["market"],
            axis=1,
        )
        fig, ax = _report_figure((9, max(4.8, 1.8 + len(data) * .38)))
        market_labels = data.market.map(lambda value: _wrap_report_label(value, width=10))
        ax.barh(market_labels, data.gmv, color=COLORS["teal"])
        _style_report_axis(ax, "市场GMV（成交总额）排名", xlabel="GMV（成交总额）", grid_axis="x")
        ax.spines["left"].set_visible(False)
        path = directory / "region_gmv.png"
        _save_report_chart(fig, path)
        charts["region"] = path

    customer = _result(bundle, "customer")
    if customer:
        data = customer.data["segments"].sort_values("customers")
        fig, ax = _report_figure((9, max(4.5, 1.8 + len(data) * .48)))
        segment_labels = data.segment.map(lambda value: _wrap_report_label(value, width=10))
        ax.barh(segment_labels, data.customers, color=[COLORS["grid"], COLORS["amber"], COLORS["coral"], COLORS["teal"]][:len(data)])
        _style_report_axis(ax, "客户价值分群", xlabel="客户数", grid_axis="x")
        ax.spines["left"].set_visible(False)
        path = directory / "customer_segments.png"
        _save_report_chart(fig, path)
        charts["customer"] = path
    return charts


def _report_scope_label(report_scope: str, scope_value: Any) -> str:
    labels = {"overall": "整体", "market": "指定市场", "category": "指定品类"}
    base = labels.get(report_scope, "整体")
    return "{} · {}".format(base, scope_value) if report_scope != "overall" and scope_value else base


def _append_business_docx_sections(report, business_analysis, start_index: int) -> int:
    labels = {"advertising": "广告分析", "returns": "退款分析", "logistics": "物流分析"}
    index = start_index
    for topic in ("advertising", "returns", "logistics"):
        payload = (business_analysis or {}).get(topic)
        if not payload:
            continue
        report.heading("{}. {}（模拟数据）".format(index, labels[topic]), page_break=True)
        source = payload.get("data_source", {})
        report.paragraph(
            "数据来源：{}。{}".format(
                source.get("description") or source.get("label") or "synthetic_extension",
                "本节仅用于功能验证，不代表真实经营表现。",
            )
        )
        report.heading("指标与口径", level=2)
        metrics = [
            "{}：{}{}；公式：{}；证据：{}。".format(
                item.get("label"), _clip_report_text(item.get("value"), 32),
                " {}".format(item["currency"]) if item.get("currency") else "",
                item.get("formula"), item.get("evidence_id"),
            )
            for item in payload.get("metrics", [])
        ]
        if metrics:
            report.bullets(metrics)
        else:
            report.paragraph("当前分析范围没有可计算的注册指标。")
        report.heading("异常、阈值与证据", level=2)
        anomalies = [
            "{}（{}）；规则 {} v{}；阈值：{}；原因：{}；证据：{}。".format(
                item.get("title"), item.get("entity"), item.get("rule_id"),
                item.get("rule_version"), item.get("threshold"), item.get("reason"),
                _clip_report_text(item.get("evidence_ids"), 120),
            )
            for item in payload.get("anomalies", [])
        ]
        if anomalies:
            report.bullets(anomalies)
        else:
            report.paragraph("当前分析范围未命中版本化异常规则。")
        report.heading("行动建议", level=2)
        actions = [
            "{}：{}；复核阈值：{}；证据：{}。".format(
                item.get("priority"), item.get("title"), item.get("threshold"),
                _clip_report_text(item.get("evidence_ids"), 120),
            )
            for item in payload.get("actions", [])
        ]
        if actions:
            report.bullets(actions)
        else:
            report.paragraph("当前证据未达到行动阈值，建议维持监测。")
        limitations = list(dict.fromkeys(
            payload.get("limitations", []) + payload.get("quality", {}).get("limitations", [])
        ))
        report.heading("数据限制", level=2)
        report.bullets(limitations or ["模拟扩展数据不用于真实经营决策。"])
        index += 1
    return index


def _summary_docx(
    bundle, output_path: Path, report_scope="overall", scope_value=None,
    include_action_details=True, business_analysis=None,
) -> Path:
    currency = bundle.context.metadata.get("target_currency", "")
    report = DocxReportBuilder("跨境电商经营分析与行动报告")
    report.cover("跨境电商经营分析与行动报告 · 摘要版", {
        "报告范围": _report_scope_label(report_scope, scope_value),
        "数据文件": bundle.context.metadata.get("filename", ""),
        "生成时间": bundle.generated_at,
        "基准币种": currency_label(currency),
    })
    artifacts = getattr(bundle, "artifacts", None)
    markets, products, actions = _scope_frames(artifacts, report_scope, scope_value)
    overview = _result(bundle, "overview")
    report.heading("1. 本期经营结论")
    if overview:
        data = overview.data
        report.bullets([
            "GMV（成交总额）{}，订单 {:,}，AOV（平均客单价）{}。".format(_money(data["gmv"], currency), data["orders"], _money(data["aov"], currency)),
            "利润率 {}，退货率 {}。".format(_pct(data.get("profit_rate")), _pct(data.get("return_rate"))),
            "本期识别健康增长市场 {} 个、风险增长市场 {} 个、需优先处理商品 {} 个。".format(
                int(markets.status.eq("健康增长").sum()) if not markets.empty else 0,
                int(markets.status.eq("风险增长").sum()) if not markets.empty else 0,
                int(products.opportunity_type.isin(["高风险增长", "利润修复机会"]).sum()) if not products.empty else 0,
            ),
        ])
    report.heading("2. 主要增长来源", page_break=True)
    if not markets.empty:
        report.table(localize_frame(_flat_frame(markets[[column for column in (
            "market", "status", "current_gmv", "gmv_change", "gmv_growth_rate", "primary_driver", "evidence_ids"
        ) if column in markets]].head(8))), 8)
    else:
        report.paragraph("当前范围没有达到展示门槛的市场增长对象。")
    report.heading("3. 最重要的产品机会")
    if artifacts:
        report.paragraph(
            "跨两期合计少于 {} 单的 {:,} 个商品仅计入样本不足总数，不生成确定性机会或行动项。".format(
                artifacts.opportunity_summary.get("product_candidate_min_orders", 3),
                artifacts.opportunity_summary.get("excluded_low_sample_products", 0),
            )
        )
    if not products.empty:
        report.table(localize_frame(_flat_frame(products[[column for column in (
                "product_id", "product_name", "opportunity_type", "target_market", "primary_customer_group", "current_gmv", "rationale", "evidence_ids"
        ) if column in products]].head(8))), 8)
    else:
        report.paragraph("当前范围没有可用的产品机会对象。")
    report.heading("4. 主要经营风险", page_break=True)
    risk_markets = markets.loc[markets.status.isin(["风险增长", "收缩待修复"])] if not markets.empty else markets
    risk_products = products.loc[products.opportunity_type.isin(["高风险增长", "利润修复机会"])] if not products.empty else products
    risks = []
    risks.extend("市场 {}：{}。".format(row.market, row.status) for row in risk_markets.head(5).itertuples())
    risks.extend("商品 {}：{}。".format(row.product_name, row.opportunity_type) for row in risk_products.head(5).itertuples())
    report.bullets(risks or ["当前未识别达到规则门槛的重点经营风险。"])
    report.heading("5. 本期优先行动")
    if include_action_details and not actions.empty:
        report.table(localize_frame(_flat_frame(actions[[column for column in (
            "priority", "entity_name", "current_judgement", "impact_amount", "evidence_confidence", "action", "guardrail_metrics", "validation_period", "stop_condition", "evidence_ids"
        ) if column in actions]].head(10))), 10)
    else:
        report.paragraph("行动明细未附加；机会与风险结论仍来自同一固定分析对象。")
    report.heading("6. 当前不能判断的问题")
    unsupported = pd.DataFrame(artifacts.unsupported_conclusions) if artifacts else pd.DataFrame()
    if not unsupported.empty:
        report.table(localize_frame(_flat_frame(unsupported.head(10))), 10)
    else:
        report.paragraph("当前关键分析能力未发现字段级阻断。")
    report.paragraph("报告中的变化与贡献不代表广告、活动、库存或竞争变化等真实经营因果。")
    _append_business_docx_sections(report, business_analysis, 7)
    return report.save(output_path)


def build_docx(
    bundle, output_path: Path, report_version="full", report_scope="overall",
    scope_value=None, include_action_details=True, business_analysis=None,
) -> Path:
    if report_version == "summary":
        return _summary_docx(
            bundle, output_path, report_scope, scope_value, include_action_details,
            business_analysis,
        )
    currency = bundle.context.metadata.get("target_currency", "")
    report = DocxReportBuilder("跨境电商经营分析与行动报告")
    report.cover("跨境电商经营分析与行动报告 · 完整版", {
        "报告范围": _report_scope_label(report_scope, scope_value),
        "数据文件": bundle.context.metadata.get("filename", ""),
        "生成时间": bundle.generated_at,
        "基准币种": currency_label(currency),
    })
    overview = _result(bundle, "overview")
    sales = _result(bundle, "sales")
    product = _result(bundle, "product")
    customer = _result(bundle, "customer")
    region = _result(bundle, "region")
    returns = _result(bundle, "returns")
    with TemporaryDirectory() as temp:
        charts = _save_charts(bundle, Path(temp))
        report.heading("1. 执行摘要")
        if overview:
            data = overview.data
            report.bullets([
                "GMV（成交总额）{}，订单 {:,}，AOV（平均客单价）{}。".format(_money(data["gmv"], currency), data["orders"], _money(data["aov"], currency)),
                "利润 {}，利润率 {}。".format(_money(data.get("profit_amount"), currency), _pct(data.get("profit_rate"))),
                "退货率 {}；退货关联GMV（成交总额）不等于实际退款损失。".format(_pct(data.get("return_rate"))),
            ])
        report.heading("2. 数据概览", page_break=True)
        report.table(pd.DataFrame([
            ["来源文件", bundle.context.metadata.get("filename", "")],
            ["原始行数", bundle.context.metadata.get("rows", "")],
            ["分析行数", len(bundle.context.analysis_data)],
            ["源币种", currency_label(bundle.context.metadata.get("source_currency", ""))],
            ["基准币种", currency_label(currency)],
            ["汇率来源", bundle.context.metadata.get("fx_provider", "")],
            ["原始哈希", bundle.context.metadata.get("sha256", "")],
        ], columns=["项目", "内容"]), max_rows=20)
        report.heading("3. 销售分析", page_break=True)
        if sales:
            report.paragraph("最近完整月 {} 的GMV（成交总额）环比 {}。不完整月份不参与结论。".format(sales.data["latest_complete_month"], _pct(sales.data["mom"])))
            if "sales" in charts: report.image(charts["sales"])
            report.table(localize_frame(sales.data["monthly"].tail(12).round(4)), 12)
        else: report.paragraph(_module_omission(bundle, "sales"))
        report.heading("4. 商品分析", page_break=True)
        if product:
            report.paragraph("达到三笔订单门槛的商品共有 {:,} 个；低样本商品不输出淘汰建议。".format(product.data["eligible_products"]))
            if "product" in charts: report.image(charts["product"])
            report.table(localize_frame(product.data["products"].head(15).round(4)), 15)
        else: report.paragraph(_module_omission(bundle, "product"))
        report.heading("5. 客户分析", page_break=True)
        if customer:
            report.paragraph("RFM（客户价值模型）锚点日期为 {}。".format(customer.data["anchor_date"]))
            if "customer" in charts: report.image(charts["customer"])
            report.table(localize_frame(customer.data["segments"].round(2)), 10)
        else: report.paragraph(_module_omission(bundle, "customer"))
        report.heading("6. 市场分析", page_break=True)
        if region:
            if "region" in charts: report.image(charts["region"])
            report.table(localize_frame(region.data.head(15).round(4)), 15)
        else: report.paragraph(_module_omission(bundle, "region"))
        report.heading("7. 退货与运营风险", page_break=True)
        if returns:
            report.paragraph("退货率 {}，退货关联GMV（成交总额）{}。该金额仅表示风险暴露，不表示真实损失。".format(
                _pct(returns.data["summary"]["return_rate"]), _money(returns.data["summary"]["returned_gmv_exposure"], currency)
            ))
            report.table(localize_frame(returns.data["category"].head(10).round(4)), 10)
            if returns.message: report.paragraph(returns.message + "。")
        else: report.paragraph(_module_omission(bundle, "returns"))
        report.heading("8. 数据与解释风险", page_break=True)
        report.bullets([
            "原始数据未被填充、截断、删除或覆盖。",
            "负利润属于业务事实，不作为异常值删除。",
            "样本不足、缺失字段和不完整月份均显式降级。",
            "没有退款额或成本字段时，不计算真实退货损失。",
        ])
        if bundle.context.issues:
            report.table(localize_frame(pd.DataFrame([issue.to_dict() for issue in bundle.context.issues])), 20)
        report.heading("9. 商业建议", page_break=True)
        if bundle.recommendations:
            report.bullets(["{}：{}（依据：{}）".format(item["title"], item["action"], ", ".join(item["evidence_ids"])) for item in bundle.recommendations])
        else:
            report.paragraph("当前证据未达到自动建议阈值，建议保持监测。")
        artifacts = getattr(bundle, "artifacts", None)
        market_opportunities, product_opportunities, action_items = _scope_frames(
            artifacts, report_scope, scope_value,
        )
        report.heading("10. 市场增长分析", page_break=True)
        if not market_opportunities.empty:
            report.table(localize_frame(_flat_frame(market_opportunities[[column for column in (
                "market", "status", "current_gmv", "gmv_change", "gmv_growth_rate", "profit_margin",
                "return_rate", "primary_driver", "recommended_action", "stop_condition", "evidence_ids"
            ) if column in market_opportunities]])), 20)
        else:
            report.paragraph("当前范围没有可用的市场增长对象。")
        report.heading("11. 产品机会分析", page_break=True)
        report.paragraph(
            "跨两期合计少于 {} 单的 {:,} 个商品仅计入样本不足总数，不生成确定性机会或行动项。".format(
                artifacts.opportunity_summary.get("product_candidate_min_orders", 3),
                artifacts.opportunity_summary.get("excluded_low_sample_products", 0),
            ) if artifacts else ""
        )
        if not product_opportunities.empty:
            report.table(localize_frame(_flat_frame(product_opportunities[[column for column in (
                "product_id", "product_name", "category", "opportunity_type", "operating_status",
                "target_market", "primary_customer_group", "current_gmv", "gmv_growth_rate", "profit_margin", "return_rate",
                "rationale", "recommended_action", "stop_condition", "evidence_ids"
            ) if column in product_opportunities]])), 30)
        else:
            report.paragraph("当前范围没有可用的产品机会对象。")
        report.heading("12. 优先行动计划", page_break=True)
        if include_action_details and not action_items.empty:
            report.table(localize_frame(_flat_frame(action_items[[column for column in (
                "priority", "source_type", "entity_name", "current_judgement", "impact_amount", "evidence_confidence", "primary_driver",
                "action", "owner_role", "guardrail_metrics", "validation_period", "stop_condition", "evidence_ids"
            ) if column in action_items]])), 40)
        else:
            report.paragraph("行动明细未附加。")
        report.heading("13. 经营异常与行动建议", page_break=True)
        if artifacts and artifacts.insights:
            top_insights = pd.DataFrame(artifacts.records("insights")).head(15)
            report.paragraph("异常、诊断和建议来自同一版本化指标与证据链；下表按影响、风险、紧急度和证据可信度排序。")
            report.table(localize_frame(_flat_frame(top_insights[[column for column in (
                "priority", "severity", "entity_name", "metric_id", "finding",
                "diagnosis_status", "recommendation_summary", "limitations",
            ) if column in top_insights]])), 15)
        else:
            report.paragraph("当前未发现达到规则阈值的异常。")
        report.heading("14. 数据质量与经营分析边界", page_break=True)
        if artifacts and artifacts.data_quality:
            quality = artifacts.data_quality
            report.paragraph("数据质量评分 {:.2f}/100，可信度等级 {}：{}。".format(
                quality.quality_score, quality.quality_rating, quality.quality_explanation
            ))
            report.table(localize_frame(pd.DataFrame(artifacts.records("data_quality_dimensions"))), 10)
            report.heading("当前无法判断的问题", level=2)
            unsupported = pd.DataFrame(artifacts.unsupported_conclusions)
            if not unsupported.empty:
                report.table(localize_frame(_flat_frame(unsupported)), 20)
            report.heading("数据质量提升建议", level=2)
            improvements = pd.DataFrame(artifacts.records("data_improvement_plan"))
            if not improvements.empty:
                report.table(localize_frame(_flat_frame(improvements)), 20)
        summary_index = _append_business_docx_sections(report, business_analysis, 15)
        report.heading("{}. 总结".format(summary_index))
        report.paragraph("本报告从经营规模、销售变化、商品、客户、区域、退货及多业务专题形成可追溯分析。所有建议均绑定证据编号，未将相关性写成因果关系。")
        report.heading("附录：SQL 与指标证据索引", level=2)
        evidence_frame = pd.DataFrame(artifacts.records("evidence")) if artifacts else bundle.evidence_frame()
        report.paragraph("本附录仅展示可审计索引；完整查询参数、指标快照关联和原始关联键保留在 Excel 与 manifest。")
        report.bullets(_docx_evidence_bullets(evidence_frame) or ["当前没有可用的 SQL 证据索引。"])
        return report.save(output_path)


def _json_safe(value):
    if isinstance(value, pd.DataFrame):
        return value.to_dict("records")
    if isinstance(value, (pd.Timestamp, datetime)):
        return value.isoformat()
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, Path):
        return str(value)
    raise TypeError("Not JSON serializable: {}".format(type(value).__name__))


def export_bundle(
    bundle, output_dir: Any, report_version="full", report_scope="overall",
    scope_value=None, include_action_details=True, business_analysis=None,
) -> Dict[str, Path]:
    directory = Path(output_dir)
    directory.mkdir(parents=True, exist_ok=True)
    excel_path = ExcelBundleWriter().write(
        output_frames(bundle, report_scope, scope_value, business_analysis), directory / "analysis_result.xlsx"
    )
    _style_excel_report(excel_path)
    paths = {
        "excel": excel_path,
        "markdown": directory / "cross_border_analysis_report.md",
        "docx": directory / "cross_border_analysis_report.docx",
        "manifest": directory / "analysis_manifest.json",
    }
    paths["markdown"].write_text(build_markdown(bundle, business_analysis), encoding="utf-8")
    build_docx(
        bundle, paths["docx"], report_version, report_scope, scope_value,
        include_action_details, business_analysis,
    )
    artifacts = getattr(bundle, "artifacts", None)
    scope_id = bundle.metadata.get("scope_id")
    manifest = {
        "project": "CrossBorder AI Analytics",
        "version": "4.0.0",
        "generated_at": bundle.generated_at,
        "dataset_id": bundle.metadata.get("dataset_id"),
        "scope_id": scope_id,
        "analysis_request": bundle.metadata.get("analysis_request", {}),
        "source": bundle.context.lineage,
        "metadata": {key: value for key, value in bundle.context.metadata.items() if key != "fx_rates"},
        "filters": bundle.metadata.get("filters", {}),
        "storage": {
            "backend": bundle.metadata.get("analysis_backend"),
            "database_schema_version": bundle.metadata.get("database_schema_version"),
            "dataset_id": bundle.metadata.get("dataset_id"),
            "query_runs": bundle.metadata.get("query_runs", []),
            "market_dimension": bundle.metadata.get("market_dimension"),
            "product_detail_queries": bundle.metadata.get("product_detail_queries", []),
        },
        "module_status": bundle.status_frame().to_dict("records"),
        "metric_definitions": artifacts.records("metric_definitions") if artifacts else [],
        "metric_snapshots": artifacts.records("metric_snapshots") if artifacts else [],
        "anomalies": artifacts.records("anomalies") if artifacts else [],
        "diagnoses": artifacts.records("diagnoses") if artifacts else [],
        "recommendations": artifacts.records("recommendations") if artifacts else bundle.recommendations,
        "evidence": artifacts.records("evidence") if artifacts else bundle.evidence_frame().to_dict("records"),
        "insights": artifacts.records("insights") if artifacts else [],
        "market_opportunities": artifacts.records("market_opportunities") if artifacts else [],
        "product_opportunities": artifacts.records("product_opportunities") if artifacts else [],
        "action_items": artifacts.records("action_items") if artifacts else [],
        "opportunity_summary": dict(artifacts.opportunity_summary) if artifacts else {},
        "report_options": {
            "version": report_version, "scope": report_scope,
            "scope_value": scope_value, "include_action_details": include_action_details,
        },
        "data_quality": artifacts.data_quality.to_dict() if artifacts and artifacts.data_quality else {},
        "field_quality": artifacts.records("field_quality") if artifacts else [],
        "analysis_capability": artifacts.records("analysis_capability") if artifacts else [],
        "data_improvement_plan": artifacts.records("data_improvement_plan") if artifacts else [],
        "unsupported_conclusions": artifacts.unsupported_conclusions if artifacts else [],
        "business_analysis": business_analysis or {},
        "legacy_evidence": bundle.evidence_frame().to_dict("records"),
        "outputs": {key: path.name for key, path in paths.items()},
    }
    paths["manifest"].write_text(json.dumps(manifest, ensure_ascii=False, indent=2, default=_json_safe), encoding="utf-8")
    return paths
