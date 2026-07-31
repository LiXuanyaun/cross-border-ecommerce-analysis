from pathlib import Path
from types import SimpleNamespace
import json

import pytest
from fastapi.testclient import TestClient
from docx import Document
from openpyxl import load_workbook

from crossborder_analytics.database import CrossBorderDatabase
from crossborder_analytics.multibusiness_analysis import MultiBusinessAnalysisService
from crossborder_analytics.multibusiness_import import MultiBusinessImportService
from crossborder_analytics.reporting import export_bundle
from crossborder_analytics.service import AnalysisService
from crossborder_api.agent_context_builder import AgentContextBuilder
from crossborder_api.agent_tools import build_agent_tool_registry
from crossborder_api.app_state import runtime
from crossborder_api.main import app
from crossborder_api.multibusiness_presenter import MultiBusinessPresenter


FIXTURE = Path(__file__).parent / "fixtures" / "adventureworks_v4"


@pytest.fixture()
def business_dataset(tmp_path, monkeypatch):
    database = tmp_path / "business.db"
    core = AnalysisService(database_path=database)
    result = MultiBusinessImportService(database, core).import_directory(FIXTURE)
    service = MultiBusinessAnalysisService(database)
    monkeypatch.setattr(runtime.multi_business_presenter, "analysis_service", service)
    return service, result["dataset_id"]


def test_three_planted_scenarios_produce_versioned_rules_and_evidence(business_dataset):
    service, dataset_id = business_dataset
    expected = {
        "advertising": "AD_SPEND_UP_CVR_DOWN",
        "returns": "RETURN_SIZE_RATE_SPIKE",
        "logistics": "LOGISTICS_CUSTOMS_DELAY_SPIKE",
    }

    for topic, rule_id in expected.items():
        payload = service.analyze(topic, dataset_id)
        anomalies = {item["rule_id"]: item for item in payload["anomalies"]}
        assert rule_id in anomalies
        anomaly = anomalies[rule_id]
        assert anomaly["rule_version"] == "1.0.0"
        assert anomaly["threshold"]
        assert anomaly["evidence_ids"]
        assert anomaly["limitations"]
        evidence = {item["evidence_id"]: item for item in payload["evidence"]}
        assert anomaly["evidence_ids"][0] in evidence
        assert evidence[anomaly["evidence_ids"][0]]["formula"]
        assert evidence[anomaly["evidence_ids"][0]]["record_keys"]
        assert payload["data_source"]["is_simulated"] is True
        assert payload["quality"]["status"] == "WARNING"


def test_advertising_metrics_use_usd_and_registered_formulas(business_dataset):
    service, dataset_id = business_dataset
    payload = service.analyze("advertising", dataset_id, country="DE", channel="Search")
    metrics = {item["id"]: item for item in payload["metrics"]}

    assert metrics["ad.spend_usd"]["value"] == pytest.approx(150.0)
    assert metrics["ad.roas"]["value"] == pytest.approx(220 / 150)
    assert metrics["ad.spend_usd"]["currency"] == "USD"
    assert "attributed_revenue_usd" in metrics["ad.roas"]["formula"]
    assert payload["filters"] == {"country": "DE", "channel": "Search"}


def test_out_of_range_and_empty_business_facts_do_not_create_zero_rankings(business_dataset):
    service, dataset_id = business_dataset

    out_of_range = service.analyze("returns", dataset_id, start="2030-01-01", end="2030-12-31")
    filtered_empty = service.analyze("advertising", dataset_id, campaign_id="missing-campaign")

    for payload, state in ((out_of_range, "OUT_OF_RANGE"), (filtered_empty, "EMPTY")):
        assert payload["data_state"] == state
        assert all(item["value"] is None for item in payload["metrics"])
        assert payload["trend"]["rows"] == []
        assert payload["ranking"]["rows"] == []
        assert payload["details"] == []
        assert payload["anomalies"] == []


def test_business_api_contract_exposes_scope_filters_evidence_and_simulation(business_dataset):
    service, dataset_id = business_dataset
    client = TestClient(app)

    response = client.get(
        "/api/v1/business/returns",
        params={"dataset_id": dataset_id, "country": "Germany", "category": "Clothing"},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "SUCCESS"
    assert body["meta"]["dataset_id"] == dataset_id
    assert body["meta"]["scope_id"] == body["data"]["scope_id"]
    assert body["data"]["topic"] == "returns"
    assert body["data"]["filter_options"]["category"] == ["Clothing"]
    assert body["data"]["metrics"]
    assert body["data"]["evidence"]
    assert body["data"]["pagination"]["page"] == 1
    assert body["data"]["pagination"]["total"] >= len(body["data"]["details"])
    contracts = {item["type"]: item for item in body["data"]["visualizations"]}
    assert contracts["line"]["rows"] == body["data"]["trend"]["rows"]
    assert contracts["bar"]["rows"] == body["data"]["ranking"]["rows"]
    assert body["data"]["table"]["rows"] == body["data"]["details"]
    assert body["data"]["table"]["pagination"] == body["data"]["pagination"]
    assert body["data"]["data_source"]["label"] == "模拟数据"
    assert any("模拟数据" in item for item in body["limitations"])


def test_business_api_rejects_unknown_topic(business_dataset):
    _, dataset_id = business_dataset
    response = TestClient(app).get("/api/v1/business/inventory", params={"dataset_id": dataset_id})
    assert response.status_code == 400


def test_business_presenter_reuses_analysis_across_pages_and_can_invalidate():
    class AnalysisStub:
        calls = 0

        def analyze(self, topic, dataset_id, **filters):
            self.calls += 1
            return {"topic": topic, "dataset_id": dataset_id, "details": [{"id": value} for value in range(30)]}

    service = AnalysisStub()
    presenter = MultiBusinessPresenter(service)

    first = presenter.present("advertising", "dataset-1", page=1, page_size=10, country="DE")
    second = presenter.present("advertising", "dataset-1", page=2, page_size=10, country="DE")
    assert service.calls == 1
    assert first["details"][0]["id"] == 0
    assert second["details"][0]["id"] == 10

    presenter.clear_cache()
    presenter.present("advertising", "dataset-1", page=1, page_size=10, country="DE")
    assert service.calls == 2


def test_reports_share_business_metrics_evidence_and_simulation_disclosure(tmp_path):
    database = tmp_path / "business-report.db"
    core = AnalysisService(database_path=database)
    imported = MultiBusinessImportService(database, core).import_directory(FIXTURE)
    dataset_id = imported["dataset_id"]
    context = CrossBorderDatabase(database).load_context(dataset_id)
    assert context is not None
    bundle = core.run(context)
    service = MultiBusinessAnalysisService(database)
    analyses = {topic: service.analyze(topic, dataset_id) for topic in ("advertising", "returns", "logistics")}

    paths = export_bundle(bundle, tmp_path / "report", business_analysis=analyses)
    workbook = load_workbook(paths["excel"], read_only=True)
    assert {"广告分析", "退款分析", "物流分析", "多业务证据"}.issubset(workbook.sheetnames)
    workbook.close()
    styled_workbook = load_workbook(paths["excel"], read_only=False)
    advertising_sheet = styled_workbook["广告分析"]
    assert advertising_sheet["A1"].fill.fgColor.rgb == "0017212B"
    assert advertising_sheet["A2"].fill.fgColor.rgb == "00F7F8F6"
    assert advertising_sheet.sheet_view.showGridLines is False
    styled_workbook.close()
    markdown = paths["markdown"].read_text(encoding="utf-8")
    assert "广告分析（模拟数据）" in markdown
    assert "synthetic_extension" in markdown
    headings = [p.text for p in Document(paths["docx"]).paragraphs if p.style.name.startswith("Heading")]
    assert any("广告分析（模拟数据）" in item for item in headings)
    assert any("退款分析（模拟数据）" in item for item in headings)
    assert any("物流分析（模拟数据）" in item for item in headings)
    manifest = json.loads(paths["manifest"].read_text(encoding="utf-8"))
    assert manifest["version"] == "4.0.0"
    assert manifest["dataset_id"] == dataset_id
    assert set(manifest["business_analysis"]) == {"advertising", "returns", "logistics"}
    assert manifest["business_analysis"]["advertising"]["metrics"] == analyses["advertising"]["metrics"]
    assert manifest["business_analysis"]["returns"]["evidence"] == analyses["returns"]["evidence"]


def test_agent_context_uses_registered_business_tools_without_arbitrary_sql(business_dataset):
    service, dataset_id = business_dataset
    database = service.database_path
    core = AnalysisService(database_path=database)
    context = CrossBorderDatabase(database).load_context(dataset_id)
    assert context is not None
    bundle = core.run(context)

    class RuntimeStub:
        multi_business_analysis_service = service
        agent_tools = build_agent_tool_registry()

        @staticmethod
        def overview(_dataset_id, _start, _end):
            return {"period": {"start": "2013-05-01", "end": "2013-10-01"}}, bundle

        @staticmethod
        def bundle(_dataset_id, _start, _end, _market, _category):
            return bundle

        @staticmethod
        def scenario(_dataset_id):
            return SimpleNamespace(name="AdventureWorks 多业务分析")

    agent_context, _ = AgentContextBuilder(RuntimeStub()).build(
        dataset_id, "请解释广告、退款和物流异常并给出行动建议"
    )
    assert set(agent_context["business_analysis"]) == {"advertising", "returns", "logistics"}
    assert {
        "query_business_metrics", "list_business_anomalies", "get_business_evidence",
    }.issubset(agent_context["tools"])
    observations = {item["tool"]: item for item in agent_context["tool_observations"]}
    assert observations["query_business_metrics"]["payload"]["metrics"]
    assert observations["list_business_anomalies"]["payload"]["anomalies"]
    assert observations["get_business_evidence"]["payload"]["evidence"]
    assert not any("sql" in tool.lower() for tool in agent_context["tools"])


def test_agent_context_does_not_run_unrequested_business_topics(business_dataset, monkeypatch):
    service, dataset_id = business_dataset
    database = service.database_path
    core = AnalysisService(database_path=database)
    context = CrossBorderDatabase(database).load_context(dataset_id)
    bundle = core.run(context)
    calls = []
    original = service.analyze

    def tracked_analyze(topic, selected_dataset_id, **filters):
        calls.append(topic)
        return original(topic, selected_dataset_id, **filters)

    monkeypatch.setattr(service, "analyze", tracked_analyze)

    class RuntimeStub:
        multi_business_analysis_service = service
        agent_tools = build_agent_tool_registry()

        @staticmethod
        def overview(_dataset_id, _start, _end):
            return {"period": {"start": "2013-05-01", "end": "2013-10-01"}}, bundle

        @staticmethod
        def scenario(_dataset_id):
            return SimpleNamespace(name="AdventureWorks 多业务分析")

    context_payload, _ = AgentContextBuilder(RuntimeStub()).build(dataset_id, "请分析广告投放效率")

    assert calls == ["advertising"]
    assert set(context_payload["business_analysis"]) == {"advertising"}
