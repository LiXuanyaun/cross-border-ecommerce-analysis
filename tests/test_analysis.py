from pathlib import Path
import json

import pandas as pd
from docx import Document
from openpyxl import load_workbook

from autoclean.analytics import AnalysisStatus
from crossborder_analytics.reporting import export_bundle
from crossborder_analytics.service import AnalysisService


ROOT = Path(__file__).resolve().parents[1]
SAMPLE = ROOT / "data" / "ecommerce_sales_34500.csv"


def sample_bundle():
    service = AnalysisService(cache_dir=ROOT / ".cache" / "test_fx")
    context = service.prepare(SAMPLE, source_currency="CNY", target_currency="CNY")
    return service.run(context)


def test_sample_acceptance_metrics_and_period_logic():
    bundle = sample_bundle()
    overview = bundle.results["overview"]
    assert overview.status == AnalysisStatus.SUCCESS
    data = overview.data
    assert data["gmv"] == 5865293.05
    assert data["orders"] == 34500
    assert data["customers"] == 7903
    assert data["units"] == 51430
    assert round(data["aov"], 2) == 170.01
    assert data["profit_amount"] == 970019.41
    assert round(data["profit_rate"] * 100, 4) == 16.5383
    assert round(data["return_rate"] * 100, 4) == 5.5159
    sales = bundle.results["sales"].data
    assert sales["latest_complete_month"] == "2025-08"
    assert round(sales["mom"] * 100, 4) == -2.3178
    assert "2025-09" in sales["incomplete_months"]
    assert sales["seasonality_available"] is False
    customer_segments = bundle.results["customer"].data["segments"].set_index("segment")["customers"].to_dict()
    assert customer_segments == {
        "VIP客户": 1196,
        "高价值客户": 1012,
        "流失风险客户": 1954,
        "普通客户": 3741,
    }


def test_missing_optional_modules_skip_without_failing_core(tmp_path):
    source = tmp_path / "minimum.csv"
    source.write_text(
        "order_id,order_date,total_amount\nA1,2025-01-01,10\nA2,2025-01-02,20\n",
        encoding="utf-8",
    )
    service = AnalysisService(cache_dir=tmp_path / "fx")
    bundle = service.run(service.prepare(source, source_currency="CNY", target_currency="CNY"))
    assert bundle.results["overview"].status == AnalysisStatus.SUCCESS
    assert bundle.results["product"].status == AnalysisStatus.SKIPPED
    assert bundle.results["customer"].status == AnalysisStatus.SKIPPED
    assert bundle.results["returns"].status == AnalysisStatus.SKIPPED


def test_duplicate_order_is_fatal(tmp_path):
    source = tmp_path / "duplicate.csv"
    source.write_text(
        "order_id,order_date,total_amount\nA1,2025-01-01,10\nA1,2025-01-02,20\n",
        encoding="utf-8",
    )
    service = AnalysisService(cache_dir=tmp_path / "fx")
    bundle = service.run(service.prepare(source, source_currency="CNY", target_currency="CNY"))
    assert bundle.results["dataset"].status == AnalysisStatus.FATAL
    assert "粒度" in bundle.results["dataset"].message


def test_uploaded_fx_converts_all_money_and_keeps_source(tmp_path):
    source = tmp_path / "fx_orders.csv"
    source.write_text(
        "order_id,order_date,total_amount,currency,profit_margin\nA1,2025-01-04,10,USD,2\n",
        encoding="utf-8",
    )
    rates = pd.DataFrame([{
        "date": "2025-01-03", "source_currency": "USD", "target_currency": "CNY", "rate": 7.2
    }])
    service = AnalysisService(cache_dir=tmp_path / "fx")
    context = service.prepare(source, target_currency="CNY", fx_rates=rates, online_fx=False)
    bundle = service.run(context)
    assert context.analysis_data.loc[0, "total_amount"] == 10
    assert context.analysis_data.loc[0, "total_amount_base"] == 72
    assert bundle.results["overview"].data["profit_amount"] == 14.4


def test_incomplete_fx_skips_consolidated_money_modules(tmp_path):
    source = tmp_path / "missing_fx.csv"
    source.write_text(
        "order_id,order_date,total_amount,currency\nA1,2025-01-04,10,USD\n",
        encoding="utf-8",
    )
    service = AnalysisService(cache_dir=tmp_path / "fx")
    context = service.prepare(source, target_currency="CNY", online_fx=False)
    bundle = service.run(context)
    assert context.metadata["fx_complete"] is False
    assert bundle.results["overview"].status == AnalysisStatus.SKIPPED
    assert bundle.results["overview"].data is None


def test_exports_have_required_surfaces_and_provenance(tmp_path):
    bundle = sample_bundle()
    source_hash = bundle.context.metadata["sha256"]
    paths = export_bundle(bundle, tmp_path)
    assert all(path.exists() for path in paths.values())
    workbook = load_workbook(paths["excel"], read_only=True)
    required = {
        "summary", "product_analysis", "customer_analysis", "region_analysis",
        "return_analysis", "module_status", "evidence", "data_quality", "fx_rates", "字段说明",
    }
    assert required.issubset(workbook.sheetnames)
    summary_headers = [cell.value for cell in next(workbook["summary"].iter_rows())]
    assert summary_headers == ["指标", "数值", "币种"]
    assert workbook["summary"]["A2"].value == "GMV（成交总额）"
    guide_headers = [cell.value for cell in next(workbook["字段说明"].iter_rows())]
    assert guide_headers == ["中文名称", "规范字段", "业务定义", "计算公式"]
    workbook.close()
    markdown = paths["markdown"].read_text(encoding="utf-8")
    assert "[overview.gmv]" in markdown
    assert "GMV（成交总额）" in markdown
    assert "RFM（客户价值模型）" in markdown
    assert "退货关联GMV（成交总额）" in markdown
    headings = [p.text for p in Document(paths["docx"]).paragraphs if p.style.name.startswith("Heading")]
    assert any("执行摘要" in text for text in headings)
    assert any("商品分析" in text for text in headings)
    manifest = json.loads(paths["manifest"].read_text(encoding="utf-8"))
    assert manifest["source"]["source_sha256"] == source_hash
    assert manifest["module_status"][0]["status"] in {"SUCCESS", "SKIPPED", "FAILED", "FATAL"}
    assert SAMPLE.read_bytes()
